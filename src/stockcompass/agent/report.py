"""Reports the agent writes: a branded HTML document (charts drawn as SVG), turned into PDF by the window,
a Word file (charts as pictures) and an Excel workbook (tables plus native Excel charts)."""

from __future__ import annotations

import html
import re
from datetime import date, datetime
from pathlib import Path
from typing import Callable

BROWN, GOLD, INK, MUTED, LINE = "#94481a", "#b8860b", "#1f1915", "#766c63", "#e7e2dc"
PALETTE = ["#94481a", "#2f74b5", "#b8860b", "#0ca30c", "#d03b3b", "#7d4b1e"]


# ------------------------------------------------------------------------------------------------ markdown
def _inline(t: str) -> str:
    t = html.escape(t)
    t = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", t)
    t = re.sub(r"(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)", r"<i>\1</i>", t)
    t = re.sub(r"`([^`]+)`", r"<code>\1</code>", t)
    return t


def md_to_html(md: str) -> str:
    lines = (md or "").replace("\r", "").split("\n")
    out, i = [], 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", ln)
        if m:
            lvl = min(4, len(m.group(1)) + 1)
            out.append(f"<h{lvl}>{_inline(m.group(2))}</h{lvl}>")
            i += 1
            continue
        if ln.lstrip().startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
            head = [c.strip() for c in ln.strip().strip("|").split("|")]
            i += 2
            body = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                body.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            out.append(table_html(head, body))
            continue
        if re.match(r"^\s*([-*•]|\d+[.)])\s+", ln):
            ordered = bool(re.match(r"^\s*\d+[.)]", ln))
            items = []
            while i < len(lines) and re.match(r"^\s*([-*•]|\d+[.)])\s+", lines[i]):
                items.append(re.sub(r"^\s*([-*•]|\d+[.)])\s+", "", lines[i]))
                i += 1
            tag = "ol" if ordered else "ul"
            out.append(f"<{tag}>" + "".join(f"<li>{_inline(x)}</li>" for x in items) + f"</{tag}>")
            continue
        para = [ln]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,4}\s|\s*([-*•]|\d+[.)])\s|\s*\|)", lines[i]):
            para.append(lines[i].rstrip())
            i += 1
        out.append("<p>" + _inline(" ".join(para)) + "</p>")
    return "\n".join(out)


def _num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def fmt_cell(v) -> str:
    if v is None:
        return "—"
    if _num(v):
        return f"{v:,.0f}" if abs(v) >= 100 or float(v).is_integer() else f"{v:,.2f}"
    return str(v)


def _numlike(c) -> bool:
    return _num(c) or bool(re.match(r"^[-+−]?[0-9.,]+\s*[%KMB]?$", str(c).replace("PKR ", "").strip()))


def table_html(head: list, rows: list[list]) -> str:
    numcol = [bool(rows) and all(_numlike(r[j]) for r in rows if j < len(r) and r[j] not in (None, "", "—")) for j in range(len(head))]
    th = "".join(f"<th class='{'n' if numcol[j] else ''}'>{_inline(str(h))}</th>" for j, h in enumerate(head))
    tr = "".join("<tr>" + "".join(f"<td class='{'n' if _numlike(c) else ''}'>{_inline(fmt_cell(c))}</td>" for c in r) + "</tr>" for r in rows)
    return f"<table><thead><tr>{th}</tr></thead><tbody>{tr}</tbody></table>"


# ------------------------------------------------------------------------------------------------ SVG charts
def _nice(lo: float, hi: float, n: int = 4):
    import math
    span = (hi - lo) or 1
    raw = span / n
    p = 10 ** math.floor(math.log10(raw))
    m = raw / p
    st = (1 if m <= 1 else 2 if m <= 2 else 2.5 if m <= 2.5 else 5 if m <= 5 else 10) * p
    return math.floor(lo / st) * st, math.ceil(hi / st) * st, st


def _fmt_axis(v: float, unit: str) -> str:
    a = abs(v)
    if unit == "pct":
        return f"{v:.0f}%"
    if a >= 1e9:
        return f"{v / 1e9:.1f}B"
    if a >= 1e6:
        return f"{v / 1e6:.0f}M" if a >= 1e8 else f"{v / 1e6:.1f}M"
    if a >= 1e3:
        return f"{v / 1e3:.0f}K"
    return f"{v:.0f}" if float(v).is_integer() else f"{v:.1f}"


def svg_chart(spec: dict, w: int = 720, h: int = 300) -> str:
    t = spec.get("type") or "bar"
    labels = [str(x) for x in spec.get("labels") or []]
    series = [s for s in spec.get("series") or [] if s.get("values")]
    unit = spec.get("unit") or "num"
    if not labels or not series:
        return ""
    font = "font-family='Segoe UI, Manrope, Arial, sans-serif'"
    if t == "hbar":
        vals = [float(v or 0) for v in series[0]["values"]][: len(labels)]
        hi = max(vals + [1])
        rh, L = 26, 190
        h = rh * len(vals) + 20
        g = []
        for i, (lab, v) in enumerate(zip(labels, vals)):
            y = 10 + i * rh
            bw = max(0, (w - L - 90) * v / hi)
            g.append(f"<text x='{L - 8}' y='{y + 16}' text-anchor='end' font-size='12' fill='{INK}' {font}>{html.escape(lab[:30])}</text>"
                     f"<rect x='{L}' y='{y + 4}' width='{w - L - 90}' height='14' rx='3' fill='#f0ece7'/>"
                     f"<rect x='{L}' y='{y + 4}' width='{bw:.1f}' height='14' rx='3' fill='{BROWN}'/>"
                     f"<text x='{L + bw + 6:.1f}' y='{y + 16}' font-size='12' font-weight='700' fill='{INK}' {font}>{_fmt_axis(v, unit)}</text>")
        return f"<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 {w} {h}' width='{w}' height='{h}'>{''.join(g)}</svg>"
    Lm, R, T, B = 60, 16, 30, 46
    allv = [float(v) for s in series for v in s["values"] if v is not None]
    lo = 0 if t == "bar" or min(allv) >= 0 else min(allv)
    lo, hi, st = _nice(lo, max(allv + [lo + 1]))
    if t == "bar":
        lo = min(lo, 0)
    y = lambda v: T + (h - T - B) * (1 - (v - lo) / ((hi - lo) or 1))
    g = []
    v = lo
    while v <= hi + 1e-9:
        g.append(f"<line x1='{Lm}' x2='{w - R}' y1='{y(v):.1f}' y2='{y(v):.1f}' stroke='{LINE}'/>"
                 f"<text x='{Lm - 8}' y='{y(v) + 4:.1f}' text-anchor='end' font-size='11' fill='{MUTED}' {font}>{_fmt_axis(v, unit)}</text>")
        v += st
    n = len(labels)
    legend = "".join(f"<rect x='{Lm + i * 150}' y='6' width='11' height='11' rx='2' fill='{PALETTE[i % 6]}'/>"
                     f"<text x='{Lm + i * 150 + 16}' y='16' font-size='12' fill='{INK}' {font}>{html.escape(s.get('name') or '')[:22]}</text>"
                     for i, s in enumerate(series))
    every = max(1, (n + 9) // 10)
    if t == "line":
        x = lambda i: Lm + (w - Lm - R) * (0.5 if n == 1 else i / (n - 1))
        for si, s in enumerate(series):
            pts = " ".join(f"{x(i):.1f},{y(float(v)):.1f}" for i, v in enumerate(s["values"][:n]) if v is not None)
            g.append(f"<polyline points='{pts}' fill='none' stroke='{PALETTE[si % 6]}' stroke-width='2.2' stroke-linejoin='round'/>")
        for i, lab in enumerate(labels):
            if i % every == 0 or i == n - 1:
                g.append(f"<text x='{x(i):.1f}' y='{h - 24}' text-anchor='middle' font-size='11' fill='{MUTED}' {font}>{html.escape(_short(lab))}</text>")
    else:
        band = (w - Lm - R) / n
        bw = min(34, band * 0.7 / len(series))
        for i, lab in enumerate(labels):
            cx = Lm + band * i + band / 2
            gw = bw * len(series) + 2 * (len(series) - 1)
            for si, s in enumerate(series):
                val = float((s["values"] + [0] * n)[i] or 0)
                top, bot = y(max(val, 0)), y(min(val, 0) if lo < 0 else lo)
                g.append(f"<rect x='{cx - gw / 2 + si * (bw + 2):.1f}' y='{top:.1f}' width='{bw:.1f}' height='{max(1, bot - top):.1f}' rx='3' fill='{PALETTE[si % 6]}'/>")
            if i % every == 0:
                g.append(f"<text x='{cx:.1f}' y='{h - 24}' text-anchor='middle' font-size='11' fill='{MUTED}' {font}>{html.escape(_short(lab))}</text>")
    return f"<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 {w} {h}' width='{w}' height='{h}'>{legend}{''.join(g)}</svg>"


def _short(lab: str) -> str:
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})$", lab)
    if m:
        return f"{int(m.group(3))} {['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun', 'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec'][int(m.group(2)) - 1]}"
    return lab[:16]


# ------------------------------------------------------------------------------------------------ documents
CSS = f"""
@page {{ size: A4; margin: 16mm 14mm; }}
body {{ font-family: 'Segoe UI', Manrope, Arial, sans-serif; color: {INK}; font-size: 11pt; line-height: 1.45; margin: 0; }}
.wrap {{ max-width: 780px; margin: 0 auto; padding: 24px; }}
.band {{ border-top: 6px solid {BROWN}; padding-top: 14px; display: flex; justify-content: space-between; align-items: flex-end; }}
.brand {{ color: {BROWN}; font-weight: 800; letter-spacing: .04em; font-size: 10pt; text-transform: uppercase; }}
h1 {{ font-size: 22pt; margin: 6px 0 2px; }} h2, h3, h4, figcaption {{ break-after: avoid; page-break-after: avoid; }} h2 {{ font-size: 14pt; color: {BROWN}; margin: 22px 0 6px; border-bottom: 1px solid {LINE}; padding-bottom: 4px; }}
h3 {{ font-size: 12pt; margin: 14px 0 4px; }} h4 {{ font-size: 11pt; margin: 10px 0 4px; }}
.sub {{ color: {MUTED}; }} .summary {{ background: #faf6ee; border-left: 4px solid {GOLD}; padding: 10px 14px; border-radius: 6px; margin: 14px 0; }}
table {{ border-collapse: collapse; width: 100%; font-size: 9.5pt; margin: 8px 0 12px; page-break-inside: avoid; }}
th {{ background: {BROWN}; color: #fff; text-align: left; padding: 5px 7px; font-weight: 700; }} th.n {{ text-align: right; }}
td {{ border-bottom: 1px solid {LINE}; padding: 4px 7px; }} td.n {{ text-align: right; font-variant-numeric: tabular-nums; }}
tr:nth-child(even) td {{ background: #faf9f7; }}
figure {{ margin: 10px 0 14px; page-break-inside: avoid; break-inside: avoid; }} figcaption {{ font-weight: 700; margin-bottom: 4px; }}
figure svg {{ width: 100%; height: auto; }}
.foot {{ color: {MUTED}; font-size: 8.5pt; margin-top: 24px; border-top: 1px solid {LINE}; padding-top: 6px; }}
code {{ background: #f3efe9; padding: 0 4px; border-radius: 3px; }}
"""


def build_html(spec: dict) -> str:
    title = spec.get("title") or "Report"
    parts = [f"<div class='band'><div><div class='brand'>Stock Compass · MAF Carrefour Pakistan</div><h1>{html.escape(title)}</h1>"
             f"<div class='sub'>{html.escape(spec.get('subtitle') or '')}</div></div><div class='sub'>{date.today():%d %b %Y}</div></div>"]
    if spec.get("summary"):
        parts.append(f"<div class='summary'><b>Summary</b>{md_to_html(spec['summary'])}</div>")
    for s in spec.get("sections") or []:
        parts.append(f"<h2>{html.escape(s.get('heading') or '')}</h2>")
        if s.get("text"):
            parts.append(md_to_html(s["text"]))
        if s.get("chart"):
            svg = svg_chart(s["chart"])
            if svg:
                parts.append(f"<figure><figcaption>{html.escape(s['chart'].get('title') or '')}</figcaption>{svg}</figure>")
        tb = s.get("table")
        if tb and tb.get("columns"):
            parts.append(table_html(tb["columns"], tb.get("rows") or []))
    parts.append(f"<div class='foot'>Prepared by the Stock Compass agent from the reports loaded on this PC. "
                 f"Generated {datetime.now():%d %b %Y %H:%M}.</div>")
    return f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(title)}</title><style>{CSS}</style></head><body><div class='wrap'>{''.join(parts)}</div></body></html>"


def build_docx(spec: dict, path: Path, svg_to_png: Callable[[str], bytes | None] | None = None):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Cm
    import io

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Segoe UI"
    st.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
    brand = doc.add_paragraph()
    r = brand.add_run("STOCK COMPASS · MAF CARREFOUR PAKISTAN")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(8.5), RGBColor(0x94, 0x48, 0x1a)
    h = doc.add_heading(spec.get("title") or "Report", level=0)
    if spec.get("subtitle"):
        doc.add_paragraph(spec["subtitle"]).runs[0].font.color.rgb = RGBColor(0x76, 0x6c, 0x63)
    doc.add_paragraph(f"{date.today():%d %B %Y}").runs[0].font.color.rgb = RGBColor(0x76, 0x6c, 0x63)

    def add_md(md: str):
        for block in re.split(r"\n\s*\n", md or ""):
            lines = [ln for ln in block.split("\n") if ln.strip()]
            if not lines:
                continue
            if lines[0].lstrip().startswith("|") and len(lines) > 1 and re.match(r"^\s*\|?\s*:?-{2,}", lines[1]):
                head = [c.strip() for c in lines[0].strip().strip("|").split("|")]
                rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines[2:]]
                add_table(head, rows)
                continue
            for ln in lines:
                m = re.match(r"^(#{1,4})\s+(.*)", ln)
                if m:
                    doc.add_heading(m.group(2), level=min(3, len(m.group(1)) + 1))
                    continue
                bullet = re.match(r"^\s*([-*•]|\d+[.)])\s+(.*)", ln)
                p = doc.add_paragraph(style=("List Number" if bullet and bullet.group(1)[0].isdigit() else "List Bullet") if bullet else None)
                text = bullet.group(2) if bullet else ln
                for piece in re.split(r"(\*\*.+?\*\*)", text):
                    if piece.startswith("**") and piece.endswith("**"):
                        p.add_run(piece[2:-2]).bold = True
                    elif piece:
                        p.add_run(piece)

    def shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear")
        sh.set(qn("w:color"), "auto")
        sh.set(qn("w:fill"), hexcolor)
        tcPr.append(sh)

    def add_table(head, rows):
        t = doc.add_table(rows=1, cols=len(head))
        t.style = "Table Grid"
        for i, hcell in enumerate(head):
            c = t.rows[0].cells[i]
            c.text = str(hcell)
            shade(c, "94481A")
            for run in c.paragraphs[0].runs:
                run.bold, run.font.color.rgb, run.font.size = True, RGBColor(255, 255, 255), Pt(9)
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row[: len(head)]):
                cells[i].text = fmt_cell(v)
                for run in cells[i].paragraphs[0].runs:
                    run.font.size = Pt(9)
                if _num(v):
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        doc.add_paragraph()

    if spec.get("summary"):
        doc.add_heading("Summary", level=1)
        add_md(spec["summary"])
    for s in spec.get("sections") or []:
        doc.add_heading(s.get("heading") or "", level=1)
        if s.get("text"):
            add_md(s["text"])
        if s.get("chart"):
            png = svg_to_png(svg_chart(s["chart"])) if svg_to_png else None
            if png:
                doc.add_paragraph(s["chart"].get("title") or "").runs[0].bold = True
                doc.add_picture(io.BytesIO(png), width=Cm(16.5))
            else:     # no picture renderer: put the chart's numbers in a table
                ch = s["chart"]
                add_table([""] + [x.get("name") or "" for x in ch.get("series") or []],
                          [[lab] + [(x.get("values") or [None] * 999)[i] for x in ch.get("series") or []] for i, lab in enumerate(ch.get("labels") or [])])
        tb = s.get("table")
        if tb and tb.get("columns"):
            add_table(tb["columns"], tb.get("rows") or [])
    foot = doc.add_paragraph(f"Prepared by the Stock Compass agent from the reports loaded on this PC · {datetime.now():%d %b %Y %H:%M}")
    foot.runs[0].font.size, foot.runs[0].font.color.rgb = Pt(8), RGBColor(0x76, 0x6c, 0x63)
    doc.save(str(path))


def build_xlsx(spec: dict, path: Path):
    import xlsxwriter

    wb = xlsxwriter.Workbook(str(path))
    head = wb.add_format({"bold": True, "bg_color": "#94481A", "font_color": "white", "border": 1})
    title = wb.add_format({"bold": True, "font_size": 15, "font_color": "#94481A"})
    wrap = wb.add_format({"text_wrap": True, "valign": "top"})
    num = wb.add_format({"num_format": "#,##0.##"})
    ws = wb.add_worksheet("Summary")
    ws.write(0, 0, spec.get("title") or "Report", title)
    ws.write(1, 0, f"{spec.get('subtitle') or ''}  ·  {date.today():%d %b %Y}")
    ws.set_column(0, 0, 110, wrap)
    row = 3
    for chunk in [spec.get("summary") or ""] + [f"{s.get('heading')}\n{s.get('text') or ''}" for s in spec.get("sections") or []]:
        if chunk.strip():
            ws.write(row, 0, re.sub(r"\*\*|__|#+ ", "", chunk.strip()))
            row += 2
    used = set()
    for n, s in enumerate(spec.get("sections") or [], 1):
        name = re.sub(r"[\[\]:*?/\\]", "", (s.get("heading") or f"Section {n}"))[:28] or f"Section {n}"
        while name in used:
            name = name[:25] + f" {n}"
        used.add(name)
        tb, ch = s.get("table"), s.get("chart")
        if not ((tb and tb.get("columns")) or ch):
            continue
        sh = wb.add_worksheet(name)
        r0 = 0
        if ch and ch.get("labels") and ch.get("series"):
            sh.write(0, 0, ch.get("title") or "", wb.add_format({"bold": True}))
            sh.write_row(1, 0, [""] + [x.get("name") or "" for x in ch["series"]], head)
            for i, lab in enumerate(ch["labels"]):
                sh.write(2 + i, 0, lab)
                for j, x in enumerate(ch["series"]):
                    v = (x.get("values") or [])[i] if i < len(x.get("values") or []) else None
                    if _num(v):
                        sh.write_number(2 + i, 1 + j, v, num)
            kind = {"line": "line", "hbar": "bar"}.get(ch.get("type"), "column")
            c = wb.add_chart({"type": kind})
            for j, x in enumerate(ch["series"]):
                c.add_series({"name": x.get("name") or "", "categories": [sh.name, 2, 0, 1 + len(ch["labels"]), 0],
                              "values": [sh.name, 2, 1 + j, 1 + len(ch["labels"]), 1 + j], "fill": {"color": PALETTE[j % 6]},
                              "line": {"color": PALETTE[j % 6]}})
            c.set_title({"name": ch.get("title") or ""})
            c.set_legend({"position": "bottom"})
            sh.insert_chart(1, 3 + len(ch["series"]), c, {"x_scale": 1.4, "y_scale": 1.2})
            r0 = 4 + len(ch["labels"])
        if tb and tb.get("columns"):
            sh.write_row(r0, 0, tb["columns"], head)
            for i, rr in enumerate(tb.get("rows") or []):
                for j, v in enumerate(rr):
                    if _num(v):
                        sh.write_number(r0 + 1 + i, j, v, num)
                    else:
                        sh.write(r0 + 1 + i, j, "" if v is None else str(v))
            sh.autofilter(r0, 0, r0 + max(1, len(tb.get("rows") or [])), len(tb["columns"]) - 1)
        sh.set_column(0, 0, 28)
        sh.set_column(1, 12, 14)
    wb.close()


def write_report(spec: dict, folder: Path, html_to_pdf: Callable[[str, str], bool] | None = None,
                 svg_to_png: Callable[[str], bytes | None] | None = None) -> list[dict]:
    folder.mkdir(parents=True, exist_ok=True)
    safe = re.sub(r"[^\w\- ]+", "", spec.get("title") or "Report").strip()[:70] or "Report"
    base = folder / f"{date.today():%Y-%m-%d} {safe}"
    formats = [f for f in (spec.get("formats") or ["pdf", "docx", "xlsx"]) if f in ("pdf", "docx", "xlsx", "html")]
    files = []
    html_doc = build_html(spec)
    hp = base.with_suffix(".html")
    hp.write_text(html_doc, encoding="utf-8")
    if "html" in formats or "pdf" in formats:
        files.append({"kind": "html", "name": hp.name, "path": str(hp)})
    if "pdf" in formats and html_to_pdf:
        pp = base.with_suffix(".pdf")
        if html_to_pdf(str(hp), str(pp)):
            files.append({"kind": "pdf", "name": pp.name, "path": str(pp)})
    if "docx" in formats:
        dp = base.with_suffix(".docx")
        try:
            build_docx(spec, dp, svg_to_png)
            files.append({"kind": "docx", "name": dp.name, "path": str(dp)})
        except ImportError:
            pass
    if "xlsx" in formats:
        xp = base.with_suffix(".xlsx")
        build_xlsx(spec, xp)
        files.append({"kind": "xlsx", "name": xp.name, "path": str(xp)})
    return files
